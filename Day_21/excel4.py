"""wang
"""
from openpyxl import Workbook
from openpyxl.chart import BarChart,Reference

wb = Workbook(write_only=True)
sheet = wb.create_sheet()

rows=[
    ('类别', '销售A组', '销售B组'),
    ('手机', 40, 30),
    ('平板', 50, 60),
    ('笔记本', 80, 70),
    ('外围设备', 20, 10),
]
for row in rows:
    sheet.append(row)

chart =BarChart()
chart.type = "col"
chart.style = 10
chart.title ="销售统计图"
chart.y_axis.title ="销量"
chart.x_axis.title ="商品类别"
data = Reference(sheet,min_col=2,min_row=1,max_row=5,max_col=3)
cats = Reference(sheet,min_col=1,min_row=2,max_row=5)

chart.add_data(data,titles_from_data=True)
chart.set_categories(cats)
chart.shape=4
sheet.add_chart(chart,"A10")
wb.save("demo.xlsx")